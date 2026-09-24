#!/usr/bin/env python3
"""Convierte una plantilla de evaluación de ítems (.xlsx) al formato JSON del benchmark.

Lee la primera hoja de la plantilla, que debe tener una fila por ítem con las columnas
de contenido (ID, Dimension, Topic, Difficulty, Question, Expected_Answer,
Evaluation_Criteria, Source, Critical_Errors, Generator_Notes). Las columnas de revisión
(G1-G8, Q1-Q9, decisiones, notas del revisor) se ignoran: el JSON resultante contiene
solo el contenido de los ítems, conforme a data/raw/question_schema.json.

Formato esperado de las celdas multivalor:
  - Evaluation_Criteria: niveles separados por " | ", cada uno con prefijo "<n>pt: ".
  - Critical_Errors: errores separados por " | ".

Uso:
    python3 scripts/import_items_from_xlsx.py \
        --xlsx docs/new-items-evaluation/spanish_psychology_llm_new_items_evaluation.xlsx \
        --out data/raw/new_items/new_items_questions.json

Después, validar con:
    python3 scripts/validate_dataset.py --data <archivo de salida>
"""

import argparse
import json
import re
import sys
from pathlib import Path

from openpyxl import load_workbook

CRITERION_PREFIX = re.compile(r"(?:^|\s\|\s)(\d+)pt:\s*")
LIST_SEPARATOR = " | "


def parse_criteria(text: str, item_id: str) -> list[dict]:
    points = CRITERION_PREFIX.findall(text)
    parts = [p.strip() for p in CRITERION_PREFIX.split(text)[2::2]]
    if not points or len(points) != len(parts):
        raise ValueError(f"{item_id}: Evaluation_Criteria no sigue el formato '<n>pt: ... | <n>pt: ...'")
    return [{"criterion": c, "max_points": int(p)} for p, c in zip(points, parts)]


def parse_list(text):
    if text is None or not str(text).strip():
        return None
    return [e.strip() for e in str(text).split(LIST_SEPARATOR) if e.strip()]


def clean(value):
    if value is None:
        return None
    value = str(value).strip()
    return value or None


def row_to_item(row: dict) -> dict:
    item_id = clean(row["ID"])
    return {
        "id": item_id,
        "dimension": clean(row["Dimension"]),
        "topic": clean(row["Topic"]),
        "difficulty": clean(row["Difficulty"]),
        "question": clean(row["Question"]),
        "expected_answer": clean(row["Expected_Answer"]),
        "evaluation_criteria": parse_criteria(clean(row["Evaluation_Criteria"]) or "", item_id),
        "source": clean(row["Source"]),
        "notes": clean(row["Generator_Notes"]),
        "critical_errors": parse_list(row["Critical_Errors"]),
        "language": "es-ES",
    }


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--xlsx", type=Path, required=True)
    parser.add_argument("--out", type=Path, required=True)
    args = parser.parse_args()

    sheet = load_workbook(args.xlsx, read_only=True, data_only=True).worksheets[0]
    rows = sheet.iter_rows(values_only=True)
    header = next(rows)
    items = [row_to_item(dict(zip(header, r))) for r in rows if r and r[0]]

    args.out.parent.mkdir(parents=True, exist_ok=True)
    args.out.write_text(json.dumps(items, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"{len(items)} ítems escritos en {args.out}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
